# -*- coding: utf-8 -*-
"""Tests for the hz command-line tool. Standard library only; no model needed.

    python3 -m unittest discover -s tests -v

A fake completion server runs in a thread. It checks every prompt byte for byte, and "rewrites" a
draft by reversing each run of letters (digits, punctuation and links stay), so a correct rewrite has
copy rate 0 and keeps every number. Rules can make it misbehave to exercise the checks and retries.
The .docx test runs only when python-docx is installed.
"""
import contextlib
import hashlib
import io
import json
import os
import re
import socket
import subprocess
import sys
import tempfile
import threading
import unittest
from collections import Counter
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from unittest import mock

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from humanizer import hz, hz_text as T, promptfmt  # noqa: E402


def reverse_words(draft):
    parts = re.split(r'(https?://\S+)', draft)
    return ''.join(p if p.startswith('http') else re.sub(r'[A-Za-z一-鿿]+', lambda m: m.group(0)[::-1], p)
                   for p in parts)


class Fake:
    """mode 'server': llama-server (/health, /props, /tokenize, /completion).
    mode 'app': the Humanizer app (/app/status, /api/tokenize, /api/completion; POSTs need X-Humanizer: 1)."""

    def __init__(self, mode='server', rule=None, phases=None):
        self.mode, self.requests, self.calls = mode, [], Counter()
        self.rule = rule or (lambda draft, n: reverse_words(draft))
        self.phases = list(phases or [])           # app phases to report before 'ready'
        fake = self

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def _send(self, code, obj):
                raw = json.dumps(obj).encode()
                self.send_response(code)
                self.send_header('Content-Type', 'application/json')
                self.send_header('Content-Length', str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)

            def do_GET(self):
                if fake.mode == 'app' and self.path == '/app/status':
                    phase = fake.phases.pop(0) if fake.phases else 'ready'
                    return self._send(200, {'app': 'humanizer', 'phase': phase, 'tier': 'q8'})
                if fake.mode == 'server' and self.path == '/health':
                    return self._send(200, {'status': 'ok'})
                if fake.mode == 'server' and self.path == '/props':
                    return self._send(200, {'model_path': '/models/humanizer-12b-Q8_0.gguf'})
                self._send(404, {'error': 'not found'})

            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers.get('Content-Length', 0))) or b'{}')
                fake.requests.append((self.path, dict(self.headers), body))
                prefix = '/api' if fake.mode == 'app' else ''
                if fake.mode == 'app' and self.headers.get('X-Humanizer') != '1':
                    return self._send(403, {'error': 'missing X-Humanizer header'})
                if self.path == prefix + '/tokenize':
                    return self._send(200, {'tokens': list(range(len(body['content'].split())))})
                if self.path == prefix + '/completion':
                    p = body['prompt']
                    head = hz.INSTR + '\n\n'
                    assert p.startswith(head) and p.endswith(hz.SEP), 'bad prompt shape'
                    draft = p[len(head):-len(hz.SEP)]
                    assert p == promptfmt.build_prompt(draft), 'prompt differs from promptfmt.build_prompt'
                    n = fake.calls[draft]
                    fake.calls[draft] += 1
                    return self._send(200, {'content': '\n' + fake.rule(draft, n) + '\n', 'stop_type': 'eos'})
                self._send(404, {'error': 'not found'})

        self.httpd = ThreadingHTTPServer(('127.0.0.1', 0), H)
        self.url = 'http://127.0.0.1:%d' % self.httpd.server_address[1]
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()

    def completions(self):
        return [r for r in self.requests if r[0].endswith('/completion')]

    def close(self):
        self.httpd.shutdown()
        self.httpd.server_close()


def closed_url():
    s = socket.socket()
    s.bind(('127.0.0.1', 0))
    port = s.getsockname()[1]
    s.close()
    return 'http://127.0.0.1:%d' % port


def run_hz(*argv, stdin=None):
    """Run hz in-process → (exit code, stdout, stderr)."""
    o, e = io.StringIO(), io.StringIO()
    patches = [contextlib.redirect_stdout(o), contextlib.redirect_stderr(e)]
    if stdin is not None:
        patches.append(mock.patch.object(sys, 'stdin', io.TextIOWrapper(io.BytesIO(stdin))))
    with contextlib.ExitStack() as st:
        for p in patches:
            st.enter_context(p)
        code = hz.main(list(argv))
    return code, o.getvalue(), e.getvalue()


SENT = ('The team reviewed {n} support tickets and found that most of them came from the same onboarding step, '
        'which suggests the documentation could be clearer about account setup.')


def para(k, n_sent=4, base=0):
    return ' '.join(SENT.format(n=base + k * 10 + i) for i in range(n_sent))


DOC = """---
title: Quarterly review
---

# Quarterly review

{p1}

{p2}

## Numbers

Revenue grew to 1,250 dollars, up 3.5% from the 2024 baseline, according to https://example.com/report.

Key points:
- **Retention:** Customer retention improved across every region we operate in, with the largest gains in Europe.
- Short item
- Another long item explaining that the support backlog fell by 40 tickets after the new triage rota started.

```python
def total(xs):

    # a comment inside code, after a blank line
    return sum(xs)
```

| Metric | Q1 | Q2 |
|---|---|---|
| Users | 1,200 | 1,450 |

![chart](img/chart.png)

<div align="center">
  <b>Internal only</b>
</div>

$$
E = mc^2
$$

> "We will ship in March," the lead said, and that quote must stay.

**Takeaways**

{p3}

## References

1. Smith, J. (2020). A study of support tickets. Journal of Help Desks, 12(3), 45-67.
2. Doe, A. (2021). Onboarding at scale.
"""


def make_doc():
    return DOC.format(p1=para(1), p2=para(2), p3=para(3, 2))


KEPT_LINES = ['# Quarterly review', '## Numbers', 'Key points:', '- Short item', '```python', 'def total(xs):',
              '    # a comment inside code, after a blank line', '| Users | 1,200 | 1,450 |', '![chart](img/chart.png)',
              '<div align="center">', '  <b>Internal only</b>', 'E = mc^2',
              '> "We will ship in March," the lead said, and that quote must stay.', '**Takeaways**', '## References',
              '1. Smith, J. (2020). A study of support tickets. Journal of Help Desks, 12(3), 45-67.', 'title: Quarterly review']


class TestPrompt(unittest.TestCase):
    def test_fingerprint(self):
        self.assertEqual(hz.prompt_fingerprint(), 'cc51d66b4c593fbe')
        self.assertEqual(hz.FINGERPRINT, promptfmt.fingerprint())
        self.assertEqual(hashlib.sha256(hz.build_prompt('X').encode()).hexdigest()[:16], 'cc51d66b4c593fbe')

    def test_same_as_promptfmt(self):
        self.assertEqual(hz.INSTR, promptfmt.INSTR)
        self.assertEqual(hz.SEP, promptfmt.SEP)
        for d in ['X', '  draft with spaces \n\n', '第一段。\n\n第二段。', '\tTabbed\r\n']:
            self.assertEqual(hz.build_prompt(d), promptfmt.build_prompt(d))
            self.assertTrue(hz.build_prompt(d).endswith('\n\n### Rewritten:\n\n'))

    def test_sampling(self):
        self.assertEqual(hz.SAMPLING, {'temperature': 1.0, 'top_p': 0.95, 'top_k': 0, 'min_p': 0, 'repeat_penalty': 1.0})


class TestChecks(unittest.TestCase):
    def test_number_normalization(self):
        self.assertEqual(T.missing_numbers('We had 1,250 users.', 'About 1250 people.'), [])
        self.assertEqual(T.missing_numbers('Up 3.5% to $4.00', 'up 3.50 percent to $4'), [])
        self.assertEqual(T.missing_numbers('On 07 May', 'On May 7'), [])
        self.assertEqual(T.missing_numbers('２０２４年', '2024 年'), [])
        self.assertEqual(T.missing_numbers('Items 1,2,3', 'items 1, 2 and 3'), [])
        self.assertEqual(T.missing_numbers('1,250 users in 2024', 'many users in 2023'), ['1,250', '2024'])

    def test_copy_rate(self):
        s = 'the quick brown fox jumps over the lazy dog again and again'
        self.assertEqual(T.copy_rate(s, s), 1.0)
        self.assertEqual(T.copy_rate(s, reverse_words(s)), 0.0)
        zh = '今天我们讨论人工智能的发展前景和面临的挑战'
        self.assertEqual(T.copy_rate(zh, zh), 1.0)
        self.assertLess(T.copy_rate(zh, '说到人工智能，前途难料，麻烦也多'), 0.5)

    def test_check_issues(self):
        c = T.check('We shipped 3 releases and fixed 12 bugs in the last sprint.', 'We shipped some releases.')
        self.assertIn('missing_numbers', c.issues)
        self.assertEqual(c.missing_numbers, ['3', '12'])
        self.assertEqual(T.check('a b c d e f', 'f e d c b a').issues, [])
        self.assertIn('truncated', T.check('x y', 'y x', truncated=True).issues)
        long_out = 'We shipped three releases. ' + 'Then something else happened that nobody asked about. ' * 4
        c = T.check('We shipped three releases last month.', long_out)
        self.assertIn('too_long', c.issues)
        self.assertTrue(c.retry)
        c = T.check('试点门店的平均坪效同比提升了18.6%，明显高于非试点门店6.2%的同期增幅。',
                    '平均坪效同比增加18.6%（试点店），高于非试点店6.2%。试点店平均同比增加18.6%，显著高于非试点店同期6.2%增长。')
        self.assertIn('repeated', c.issues)                                   # seen on the real model
        self.assertNotIn('repeated', T.check(para(1), reverse_words(para(1))).issues)  # draft repeats itself: fine
        c = T.check('We onboarded 860 services in weekly cohorts of roughly 50 teams.',
                    'We onboarded 860 services, about 50 a week, in cohorts.<p>')        # seen on the real model
        self.assertEqual((c.issues, c.retry), (['markup'], True))
        self.assertEqual(T.check('Use <b>bold</b> here, said the 3 docs.', 'The 3 docs say: use <b>bold</b>.').issues, [])
        c = T.check('Cut spend by 35% over eight years.', 'Save $1MM+ by cutting 35% over 8 years.')
        self.assertEqual((c.added_numbers, c.retry), (['1MM'], False))   # reported, not retried

    def test_scaled_and_word_numbers(self):
        self.assertEqual(T.missing_numbers('fell from $480,000 to $305,000', 'cut from $480k to $305k'), [])
        self.assertEqual(T.missing_numbers('投入1,280万元，会员31.7万人', '投入1280万元，会员317,000人'), [])
        self.assertEqual(T.missing_numbers('in 3 weeks', 'in three weeks'), [])
        self.assertEqual(T.missing_numbers('$480k', '$480'), ['$480k'.lstrip('$')])
        self.assertEqual(T.added_numbers('第四十八家门店，用了两周', '第48家门店，用了2周'), [])


class TestPlan(unittest.TestCase):
    def test_structure_kept(self):
        doc = make_doc()
        p = T.plan_text(doc)
        kinds = Counter(u.kind for u in p.units)
        self.assertEqual(kinds['item'], 2)                       # the short item is kept
        self.assertGreaterEqual(kinds['prose'], 3)
        for kind in ('heading', 'code', 'table', 'media', 'html', 'math', 'quote', 'references', 'frontmatter', 'intro'):
            self.assertIn(kind, p.kept, kind)
        texts = '\n'.join(u.text for u in p.units)
        for line in ('def total', 'Internal only', 'E = mc^2', 'Smith, J.', 'ship in March', '| Users', 'Takeaways'):
            self.assertNotIn(line, texts)
        self.assertEqual(p.render({u.id: u.text for u in p.units}), doc)   # identity rewrite gives the input back

    def test_pieces_respect_limit_and_headings(self):
        sections = []
        for s in range(3):
            sections.append('## Section %d' % s)
            sections.extend(para(k, 3, base=100 * s) for k in range(6))     # 6 × ~90 words per section
        doc = '\n\n'.join(sections) + '\n'
        p = T.plan_text(doc)
        lines = doc.split('\n')
        heading_lines = {i for i, l in enumerate(lines) if l.startswith('## ')}
        for s, e, _ in p.repl:
            self.assertFalse(heading_lines & set(range(s, e)), 'a piece crosses a heading')
        for u in p.units:
            self.assertLessEqual(T.latin_words(u.text), 350)
        self.assertEqual([u.line for u in p.units], sorted(u.line for u in p.units))
        self.assertEqual(len(p.units), 6)                          # 540 words per section → 2 balanced pieces
        self.assertEqual(p.render({u.id: u.text for u in p.units}), doc)

    def test_long_paragraph_cut_at_sentences(self):
        long_para = para(1, 30)                                    # ~900 words, one paragraph
        p = T.plan_text('# T\n\n' + long_para + '\n')
        self.assertGreaterEqual(len(p.units), 3)
        for u in p.units:
            self.assertEqual(u.kind, 'sentences')
            self.assertLessEqual(T.latin_words(u.text), 350)
            self.assertTrue(u.text.endswith('.'))
        self.assertEqual(p.render({u.id: u.text for u in p.units}), '# T\n\n' + long_para + '\n')

    def test_chinese_pieces(self):
        zh = '这是一个关于项目进展的段落，我们在第{n}周完成了数据迁移，并且修复了大部分已知问题。' * 6
        doc = '\n\n'.join(zh.format(n=i) for i in range(6))
        p = T.plan_text(doc)
        self.assertGreater(len(p.units), 1)
        for u in p.units:
            self.assertLessEqual(T.cjk_chars(u.text), 600)
        self.assertEqual(p.render({u.id: u.text for u in p.units}), doc)

    def test_plain_text_references_and_lists(self):
        doc = ('Intro sentence that is long enough to be rewritten by the model here.\n\n'
               'References\n\n[1] Smith 2020.\n\nSmith, J. 2020. Some book. Publisher.\n')
        p = T.plan_text(doc)
        self.assertEqual(len(p.units), 1)
        self.assertIn('references', p.kept)


class ServerCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.fakes = []

    def fake(self, *a, **kw):
        f = Fake(*a, **kw)
        self.fakes.append(f)
        self.addCleanup(f.close)
        return f

    def write(self, name, text):
        path = os.path.join(self.tmp.name, name)
        with open(path, 'w', encoding='utf-8', newline='') as f:
            f.write(text)
        return path


class TestEndToEnd(ServerCase):
    def test_markdown_via_llama_server(self):
        f = self.fake('server')
        doc = make_doc()
        src, dst = self.write('in.md', doc), os.path.join(self.tmp.name, 'out.md')
        code, out, err = run_hz(src, '--server', f.url, '-o', dst, '--json', '--quiet')
        self.assertEqual(code, 0, err)
        with open(dst, encoding='utf-8') as fh:
            res = fh.read()
        for line in KEPT_LINES:
            self.assertIn(line, res.split('\n'), line)
        self.assertIn('- **Retention:** remotsuC', res)               # bullet and lead-in kept, rest rewritten
        self.assertIn('eht wen egairt ator detrats.', res)
        self.assertNotIn('The team reviewed', res)
        self.assertEqual(res.count('\n\n'), doc.count('\n\n'))         # same blank lines between blocks
        j = json.loads(out)
        plan = T.plan_text(doc)
        self.assertEqual(j['summary']['pieces'], len(plan.units))
        self.assertEqual(j['summary']['flagged'], 0)
        self.assertEqual(j['backend']['kind'], 'server')
        for piece in j['pieces']:
            for key in ('copy_rate', 'missing_numbers', 'retried', 'seconds', 'line', 'flagged', 'words_in'):
                self.assertIn(key, piece)
            self.assertEqual(piece['copy_rate'], 0.0)
        # request format: llama-server /completion, exact sampling, no stop strings
        reqs = f.completions()
        self.assertEqual(len(reqs), len(plan.units))
        for path, headers, body in reqs:
            self.assertEqual(path, '/completion')
            self.assertEqual(set(body), {'prompt', 'n_predict', 'temperature', 'top_p', 'top_k', 'min_p', 'repeat_penalty'})
            self.assertNotIn('stop', body)
            self.assertEqual((body['temperature'], body['top_p'], body['top_k'], body['min_p'], body['repeat_penalty']),
                             (1.0, 0.95, 0, 0, 1.0))
            self.assertTrue(256 <= body['n_predict'] <= 2048)
        self.assertEqual([b['prompt'] for _, _, b in reqs], [hz.build_prompt(u.text) for u in plan.units])

    def test_app_request_format(self):
        f = self.fake('app')
        src = self.write('in.txt', para(1) + '\n')
        code, out, err = run_hz(src, '--app', f.url, '--quiet')
        self.assertEqual(code, 0, err)
        self.assertEqual(out.strip(), reverse_words(para(1)))
        paths = [r[0] for r in f.requests]
        self.assertIn('/api/tokenize', paths)
        self.assertIn('/api/completion', paths)
        for path, headers, body in f.requests:
            self.assertEqual(headers.get('X-Humanizer'), '1')
            self.assertEqual(headers.get('Content-Type'), 'application/json')
        self.assertNotIn('stop', f.completions()[0][2])

    def test_server_flag_pointing_at_app(self):
        f = self.fake('app')
        code, out, err = run_hz('--server', f.url, '--quiet', stdin=(para(2) + '\n').encode())
        self.assertEqual(code, 0, err)
        self.assertEqual(f.completions()[0][0], '/api/completion')

    def test_stdin_crlf_and_chinese(self):
        f = self.fake('server')
        zh = '我们在2024年第3季度完成了1,250个工单的迁移，团队一共用了6周时间，比原计划提前了两周。\r\n\r\n第二段也需要改写，里面提到了15个客户。\r\n'
        code, out, err = run_hz('--server', f.url, '--quiet', stdin=zh.encode('utf-8'))
        self.assertEqual(code, 0, err)
        self.assertIn('\r\n', out)
        self.assertIn('1,250', out)


class TestRetry(ServerCase):
    def test_missing_number_triggers_retry(self):
        def rule(draft, n):
            out = reverse_words(draft)
            return out.replace('1,250', '1,200') if n == 0 else out
        f = self.fake('server', rule)
        src = self.write('in.txt', 'Revenue grew to 1,250 dollars this quarter, which beat the plan by a wide margin.\n')
        code, out, err = run_hz(src, '--server', f.url, '--json', '--quiet')
        self.assertEqual(code, 0, err)
        j = json.loads(out)
        piece = j['pieces'][0]
        self.assertTrue(piece['retried'])
        self.assertEqual(piece['chosen'], 2)
        self.assertFalse(piece['flagged'])
        self.assertEqual(piece['attempts'][0]['missing_numbers'], ['1,250'])
        self.assertIn('1,250', j['text'])
        self.assertEqual(len(f.completions()), 2)

    def test_copy_triggers_retry(self):
        f = self.fake('server', lambda d, n: d if n == 0 else reverse_words(d))
        code, out, err = run_hz('--server', f.url, '--json', '--quiet', stdin=(para(4) + '\n').encode())
        piece = json.loads(out)['pieces'][0]
        self.assertTrue(piece['retried'])
        self.assertEqual(piece['attempts'][0]['issues'], ['copy'])
        self.assertEqual(piece['copy_rate'], 0.0)

    def test_still_flagged_after_retry(self):
        f = self.fake('server', lambda d, n: reverse_words(d).replace('1,250', 'many'))
        src = self.write('in.md', '# Title\n\nRevenue grew to 1,250 dollars this quarter, beating the plan.\n')
        code, out, err = run_hz(src, '--server', f.url)
        self.assertEqual(code, 0)
        self.assertEqual(len(f.completions()), 2)                     # one retry, then give up
        self.assertIn('check these by hand', err)
        self.assertIn('line 3 (piece 1): numbers not found in the rewrite: 1,250', err)
        code, out, err = run_hz(src, '--server', f.url, '--json', '--quiet')
        j = json.loads(out)
        self.assertEqual(j['summary']['flagged'], 1)
        self.assertTrue(j['pieces'][0]['flagged'])
        self.assertEqual(j['pieces'][0]['missing_numbers'], ['1,250'])

    def test_retries_zero(self):
        f = self.fake('server', lambda d, n: d)
        code, out, err = run_hz('--server', f.url, '--retries', '0', '--json', '--quiet', stdin=(para(5) + '\n').encode())
        self.assertEqual(len(f.completions()), 1)
        self.assertTrue(json.loads(out)['pieces'][0]['flagged'])


POST = ("The current humanizer still carries some AI style (you can probably still feel it), and it can't get past "
        "the most accurate AI detector, Pangram. In my internal research, though, a brand-new method can actually teach "
        "a model to write like a human, not just rewrite, and it already passes Pangram v4 as a side proof.")
POST_ZH = "现在的 humanizer 还是带点 AI 味（你大概还能感觉出来），也过不了最准的 AI 检测器 Pangram。不过在我内部的研究里，有个全新的方法能真正教模型像人一样写。"
ZH_DRAFT = "我们在2024年第3季度完成了1,250个工单的迁移，团队一共用了6周时间，比原计划提前了两周，客户满意度也有明显提升。"


class TestLanguageGuard(ServerCase):
    def test_drift_rule(self):
        self.assertTrue(T.language_drift(POST, POST_ZH))
        self.assertFalse(T.language_drift(POST, reverse_words(POST)))
        self.assertFalse(T.language_drift(POST, reverse_words(POST) + " 夹了几个汉字"))      # a few Chinese characters are fine
        self.assertFalse(T.language_drift(POST + " 中文字", POST_ZH))                        # draft with 3+ Chinese characters: off
        self.assertFalse(T.language_drift(ZH_DRAFT, POST_ZH))
        self.assertTrue(T.language_drift(ZH_DRAFT, reverse_words(POST)))                    # Chinese draft written in English
        self.assertFalse(T.language_drift(ZH_DRAFT, "2024年Q3迁移了1,250个工单，6周干完，提前两周。"))
        self.assertEqual(T.LANG_RETRIES, 3)

    def test_wrong_language_is_resampled(self):
        f = self.fake('server', lambda d, n: POST_ZH if n < 2 else reverse_words(d))
        code, out, err = run_hz('--server', f.url, '--json', '--quiet', stdin=(POST + '\n').encode())
        self.assertEqual(code, 0, err)
        j = json.loads(out)
        piece = j['pieces'][0]
        self.assertEqual(len(f.completions()), 3)
        self.assertEqual(piece['language_resampled'], 2)
        self.assertFalse(piece['flagged'])
        self.assertFalse(piece['retried'])
        self.assertEqual(j['text'].strip(), reverse_words(POST))
        for _, _, body in f.completions():                                                   # same settings every time
            self.assertEqual((body['temperature'], body['top_p'], body['top_k']), (1.0, 0.95, 0))

    def test_gives_up_after_three_and_flags(self):
        f = self.fake('server', lambda d, n: POST_ZH)
        code, out, err = run_hz('--server', f.url, '--retries', '0', stdin=(POST + '\n').encode())
        self.assertEqual(code, 0)
        self.assertEqual(len(f.completions()), 4)                                            # 1 + 3 resamples, then the last one
        self.assertEqual(out.strip(), POST_ZH)
        self.assertIn('different language', err)

    def test_chinese_draft_not_touched(self):
        f = self.fake('server')
        code, out, err = run_hz('--server', f.url, '--json', '--quiet', stdin=(ZH_DRAFT + '\n').encode())
        self.assertEqual(code, 0, err)
        self.assertEqual(len(f.completions()), 1)
        self.assertEqual(json.loads(out)['pieces'][0]['language_resampled'], 0)


class TestDiscovery(ServerCase):
    def test_app_found_through_instance_json(self):
        f = self.fake('app')
        port = f.url.rsplit(':', 1)[1]
        with open(os.path.join(self.tmp.name, 'instance.json'), 'w') as fh:
            json.dump({'pid': 1, 'port': int(port)}, fh)
        with mock.patch.dict(os.environ, {'HUMANIZER_DATA_DIR': self.tmp.name}), \
                mock.patch.object(hz, 'launch_app', lambda out: self.fail('should not launch')):
            code, out, err = run_hz('--quiet', stdin=(para(6) + '\n').encode())
        self.assertEqual(code, 0, err)
        self.assertEqual(f.completions()[0][0], '/api/completion')

    def test_falls_back_to_llama_server(self):
        f = self.fake('server')
        with mock.patch.object(hz, 'app_urls', lambda: [closed_url()]), mock.patch.object(hz, 'SERVER_URL', f.url), \
                mock.patch.object(hz, 'launch_app', lambda out: self.fail('should not launch')):
            code, out, err = run_hz('--quiet', stdin=(para(7) + '\n').encode())
        self.assertEqual(code, 0, err)
        self.assertEqual(f.completions()[0][0], '/completion')

    def test_launches_app_and_waits(self):
        f = self.fake('app', phases=['starting', 'starting'])
        launched = []
        with mock.patch.object(hz, 'app_urls', lambda: [f.url] if launched else [closed_url()]), \
                mock.patch.object(hz, 'SERVER_URL', closed_url()), \
                mock.patch.object(hz, 'launch_app', lambda out: launched.append(1) or True):
            code, out, err = run_hz(stdin=(para(8) + '\n').encode())
        self.assertEqual(code, 0, err)
        self.assertEqual(launched, [1])
        self.assertIn('loading the model', err)

    def test_app_without_model(self):
        f = self.fake('app', phases=['setup'] * 5)
        with mock.patch.object(hz, 'app_urls', lambda: [f.url]), mock.patch.object(hz, 'SERVER_URL', closed_url()):
            code, out, err = run_hz(stdin=(para(9) + '\n').encode())
        self.assertEqual(code, 3)
        self.assertIn('has no model yet', err)

    def test_nothing_found(self):
        with mock.patch.object(hz, 'app_urls', lambda: [closed_url()]), mock.patch.object(hz, 'SERVER_URL', closed_url()), \
                mock.patch.object(hz, 'launch_app', lambda out: False):
            code, out, err = run_hz(stdin=(para(1) + '\n').encode())
        self.assertEqual(code, 3)
        self.assertIn('releases/latest', err)
        self.assertIn('llama-server --hf-repo jialinyyzz/humanizer', err)
        self.assertIn('docs/USAGE.md', err)
        code, out, err = run_hz('--server', closed_url(), stdin=(para(1) + '\n').encode())
        self.assertEqual(code, 3)
        self.assertIn('nothing is listening', err)


class TestCli(ServerCase):
    def test_dry_run_needs_no_model(self):
        src = self.write('in.md', make_doc())
        with mock.patch.object(hz, 'find_backend', lambda *a, **k: self.fail('dry run called a backend')):
            code, out, err = run_hz(src, '--dry-run', '--json')
        self.assertEqual(code, 0, err)
        j = json.loads(out)
        self.assertEqual(len(j['pieces']), len(T.plan_text(make_doc()).units))
        self.assertIn('code', j['kept'])

    def test_refuses_to_overwrite_input(self):
        src = self.write('in.md', 'hello')
        code, out, err = run_hz(src, '-o', src)
        self.assertEqual(code, 2)

    def test_old_scripts_still_run(self):
        r = subprocess.run([sys.executable, os.path.join(ROOT, 'humanizer', 'humanize.py'), '--help'],
                           capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        r = subprocess.run([sys.executable, os.path.join(ROOT, 'humanizer', 'hz.py'), '--version'],
                           capture_output=True, text=True)
        self.assertEqual(r.stdout.strip(), 'hz 0.1.0')


try:
    import docx  # noqa: F401
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False


@unittest.skipUnless(HAVE_DOCX, 'python-docx not installed')
class TestDocx(ServerCase):
    def make(self):
        import docx
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        d = docx.Document()
        d.add_heading('Quarterly report', level=1)
        p = d.add_paragraph('The support team closed ')
        p.add_run('1,250 tickets').bold = True
        p.add_run(' in the third quarter, which is the highest number since the team was formed in 2021.')
        d.add_paragraph('')
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = 'Metric', 'Value'
        t.cell(1, 0).text, t.cell(1, 1).text = 'Tickets closed in the quarter by the whole team', '1,250'
        link = d.add_paragraph('See the dashboard for the full numbers and the weekly breakdown: ')
        h = OxmlElement('w:hyperlink')
        r = OxmlElement('w:r')
        tx = OxmlElement('w:t')
        tx.text = 'dashboard'
        r.append(tx)
        h.append(r)
        h.set(qn('r:id'), 'rId99')
        link._p.append(h)
        d.add_paragraph('Short line.')
        d.add_heading('References', level=2)
        d.add_paragraph('Smith, J. (2020). A study of support tickets and how teams close them. Journal.')
        path = os.path.join(self.tmp.name, 'report.docx')
        d.save(path)
        return path

    def test_docx(self):
        import docx
        f = self.fake('server')
        src = self.make()
        dst = os.path.join(self.tmp.name, 'report.out.docx')
        code, out, err = run_hz(src, '--server', f.url, '-o', dst, '--json', '--quiet')
        self.assertEqual(code, 0, err)
        j = json.loads(out)
        self.assertEqual(j['summary']['pieces'], 1)
        self.assertEqual(j['pieces'][0]['paragraph'], 2)
        d = docx.Document(dst)
        ps = d.paragraphs
        self.assertEqual(ps[0].text, 'Quarterly report')
        body = ps[1]
        self.assertTrue(body.runs[0].text.startswith('ehT troppus maet desolc 1,250'))
        self.assertTrue(all(r.text == '' for r in body.runs[1:]))
        self.assertEqual(ps[2].text, '')
        self.assertTrue(ps[3].text.startswith('See the dashboard'))               # hyperlink paragraph untouched
        self.assertEqual(ps[4].text, 'Short line.')
        self.assertTrue(ps[6].text.startswith('Smith, J. (2020)'))                # references untouched
        self.assertEqual(d.tables[0].cell(1, 0).text, 'Tickets closed in the quarter by the whole team')

    def test_docx_default_output_name(self):
        f = self.fake('server')
        src = self.make()
        code, out, err = run_hz(src, '--server', f.url, '--quiet')
        self.assertEqual(code, 0, err)
        self.assertTrue(os.path.exists(os.path.join(self.tmp.name, 'report.hz.docx')))


if __name__ == '__main__':
    unittest.main()
